import openpyxl, re
p = r'C:\Users\eddys\Desktop\Matieres et fournisseurs (semrac).xlsx'
wb = openpyxl.load_workbook(p, data_only=True)
rows = []
for sheet in wb.sheetnames:
    ws = wb[sheet]
    for row in ws.iter_rows(values_only=True):
        code = row[0]
        if code is None:
            continue
        if isinstance(code, float) and code.is_integer():
            code = int(code)
        codes = str(code).strip()
        if not re.match(r'^\d', codes):
            continue  # vraies lignes data = code numérique
        fourn = row[16] if len(row) > 16 else None
        prix = row[18] if len(row) > 18 else None
        if fourn is None or str(fourn).strip() == '':
            continue  # pas de fournisseur dans l'Excel -> non reliable
        fourns = str(fourn).strip()
        try:
            prixv = float(str(prix).replace(',', '.')) if (prix is not None and str(prix).strip() != '') else None
        except Exception:
            prixv = None
        desig = str(row[1]).strip() if row[1] is not None else ''
        rows.append((codes, desig, fourns, prixv, sheet))

def esc(s): return "'" + str(s).replace("'", "''") + "'"
def num(v): return 'null' if v is None else repr(float(v))
vals = ',\n'.join("(%s,%s,%s,%s,%s)" % (esc(c), esc(d), esc(f), num(pr), esc(m)) for (c, d, f, pr, m) in rows)

sql = """drop table if exists _excel_mat;
create table _excel_mat (code text, designation text, fournisseur text, prix numeric, metal text);
insert into _excel_mat (code,designation,fournisseur,prix,metal) values
""" + vals + """;
insert into fournisseurs (nom, activite, familles_fourniture)
select distinct trim(e.fournisseur), 'Semrac', array['matiere']
from _excel_mat e
where not exists (select 1 from fournisseurs f where lower(trim(f.nom))=lower(trim(e.fournisseur)));
update stock s set fournisseur_id=f.id, prix_achat_ht=coalesce(e.prix, s.prix_achat_ht)
from _excel_mat e join fournisseurs f on lower(trim(f.nom))=lower(trim(e.fournisseur))
where s.reference=e.code;
insert into produits_fournisseurs (fournisseur_id, fournisseur_nom, reference, designation, prix, unite, date_prix, source_prix, categorie, statut, activite)
select s.fournisseur_id, f.nom, s.reference, s.designation, s.prix_achat_ht, coalesce(s.unite,'pce'),
 '2025-03-26','import','matiere_premiere',
 case when s.prix_achat_ht is null then 'en_attente_prix' else 'actif' end,
 coalesce(s.activite,'both')
from stock s join fournisseurs f on f.id=s.fournisseur_id
where s.fournisseur_id is not null
on conflict (fournisseur_id, reference) do update set prix=excluded.prix, fournisseur_nom=excluded.fournisseur_nom, designation=excluded.designation, date_prix=excluded.date_prix, statut=excluded.statut;
drop table _excel_mat;
select
 (select count(*) from stock where activite='Semrac' and fournisseur_id is not null) as semrac_lies,
 (select count(*) from stock where activite='Semrac') as semrac_total,
 (select count(*) from produits_fournisseurs) as pf_total,
 (select count(distinct fournisseur_id) from produits_fournisseurs) as pf_fournisseurs;
"""
open(r'E:\ERP derniere version\scripts_import\relink_semrac.sql', 'w', encoding='utf-8').write(sql)
print('rows:', len(rows), 'distinct fournisseurs:', len(set(f for (_, _, f, _, _) in rows)))
print('fournisseurs:', sorted(set(f for (_, _, f, _, _) in rows)))
