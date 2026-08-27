# -*- coding: utf-8 -*-
"""Import Actions correctives : 'Tableau de bord Sécurité AC' (AC + 5 Pourquoi inline) + 'PAC-SD'
-> table actions_correctives. Idempotent par id (AC-n / PAC-n / SD-n)."""
import openpyxl, urllib.request, urllib.error, json, sys, datetime
try: sys.stdout.reconfigure(encoding='utf-8', errors='backslashreplace')
except Exception: pass
SRC = r"C:\Users\eddys\OneDrive\Bureau\Seem semrac\dONN2ES\Copie de Tableau de bord QUALITE V2 07 05 2026 AGATHE.xlsm"
URL = 'https://vyqgrasezpyqjwvijwvv.supabase.co'
KEY = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InZ5cWdyYXNlenB5cWp3dmlqd3Z2Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzgwNTE4MTcsImV4cCI6MjA5MzYyNzgxN30.Optt1QD-Ai3Xm8mk4f8fsvKry9-L2SlTBNB9KgVG3ls'
HDRS = {'apikey': KEY, 'Authorization': 'Bearer ' + KEY}


def s(v):
    return (str(v).strip() or None) if v is not None and str(v).strip() else None


def d10(v):
    return v.strftime('%Y-%m-%d') if isinstance(v, (datetime.datetime, datetime.date)) else None


def numof(v):
    try: return int(float(v))
    except Exception: return None


def http(method, path, body=None, prefer=None):
    data = json.dumps(body, ensure_ascii=False).encode('utf-8') if body is not None else None
    h = dict(HDRS)
    if body is not None: h['Content-Type'] = 'application/json'
    if prefer: h['Prefer'] = prefer
    req = urllib.request.Request(URL + path, data=data, method=method, headers=h)
    with urllib.request.urlopen(req, timeout=60) as r:
        t = r.read().decode('utf-8'); return json.loads(t) if t else None


def parse_ac(wb):
    ws = wb['Tableau de bord Sécurité AC']
    recs = []
    for row in ws.iter_rows(min_row=9, max_col=18, values_only=True):
        num = numof(row[0])
        if num is None or not isinstance(row[1], (datetime.datetime, datetime.date)):
            continue
        pourquoi = [s(row[i]) for i in (8, 11, 14) if len(row) > i and s(row[i])]
        actions = [s(row[i]) for i in (9, 12, 15) if len(row) > i and s(row[i])]
        pilotes = [s(row[i]) for i in (10, 13) if len(row) > i and s(row[i])]
        recs.append({
            'id': 'AC-%d' % num, 'numero': num, 'type': 'AC', 'date_ouverture': d10(row[1]),
            'origine': s(row[3]), 'diffusion': s(row[4]), 'description': s(row[6]),
            'analyse_5p': pourquoi or None, 'cause_racine': (pourquoi[-1] if pourquoi else None),
            'action': '; '.join(actions) or None, 'responsable': '; '.join(pilotes) or None,
            'statut': 'en_cours',
        })
    return recs


def parse_pacsd(wb):
    ws = wb['Tableau de bord Sécurité PAC-SD']
    recs = []
    for row in ws.iter_rows(min_row=9, max_col=18, values_only=True):
        num = numof(row[0])
        if num is None or not isinstance(row[1], (datetime.datetime, datetime.date)):
            continue
        typ = (s(row[6]) or '').upper()
        typ = 'PAC' if 'PAC' in typ else ('SD' if 'SD' in typ else 'PAC')
        mesure = '; '.join([x for x in (s(row[11]), s(row[12]), s(row[13])) if x]) or None
        recs.append({
            'id': '%s-%d' % (typ, num), 'numero': num, 'type': typ, 'date_ouverture': d10(row[1]),
            'origine': s(row[3]), 'diffusion': s(row[4]), 'description': s(row[8]),
            'action': mesure, 'responsable': s(row[15]) if len(row) > 15 else None, 'statut': 'en_cours',
        })
    return recs


def main():
    apply = '--apply' in sys.argv
    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    recs = parse_ac(wb) + parse_pacsd(wb)
    from collections import Counter
    print('Actions parsées :', len(recs), '| par type :', dict(Counter(r['type'] for r in recs)))
    for r in recs[:6]:
        print('  ', r['id'], '|', r['date_ouverture'], '|', r['origine'], '|', str(r['description'])[:30], '| 5P:', (r.get('analyse_5p') or []))
    if not apply:
        print('\nDRY-RUN. --apply pour insérer.'); return
    existing = http('GET', '/rest/v1/actions_correctives?select=id')
    ex = set(x.get('id') for x in existing)
    todo = [r for r in recs if r['id'] not in ex]
    print('déjà en base :', len(ex), '| à insérer :', len(todo))
    for i in range(0, len(todo), 50):
        http('POST', '/rest/v1/actions_correctives', todo[i:i + 50], prefer='return=minimal')
    print('Inséré :', len(todo))


if __name__ == '__main__':
    main()
