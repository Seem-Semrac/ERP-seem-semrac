# -*- coding: utf-8 -*-
"""
Import Gestion des produits périssables (PRA2-D1, 1 feuille = 1 produit) -> produits_perissables
+ mouvements_perissables (sorties FIFO). Idempotent par (code_produit, date_reception).

Usage : python import_perissables.py            (dry-run)
        python import_perissables.py --apply
"""
import openpyxl, urllib.request, urllib.error, json, sys, datetime, unicodedata, re
try: sys.stdout.reconfigure(encoding='utf-8', errors='backslashreplace')
except Exception: pass
SRC = r"C:\Users\eddys\OneDrive\Bureau\Seem semrac\dONN2ES\PRA2 - D1Gestion des produits périssables .xlsx copie.xlsx"
URL = 'https://vyqgrasezpyqjwvijwvv.supabase.co'
KEY = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InZ5cWdyYXNlenB5cWp3dmlqd3Z2Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzgwNTE4MTcsImV4cCI6MjA5MzYyNzgxN30.Optt1QD-Ai3Xm8mk4f8fsvKry9-L2SlTBNB9KgVG3ls'
HDRS = {'apikey': KEY, 'Authorization': 'Bearer ' + KEY}
TODAY = datetime.date.today().isoformat()


def norm(v):
    t = unicodedata.normalize('NFKD', str(v or '')).encode('ascii', 'ignore').decode().lower()
    return ' '.join(t.split())


def s(v):
    return (str(v).strip() or None) if v is not None and str(v).strip() else None


def d10(v):
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime('%Y-%m-%d')
    return None


def numof(v):
    try: return float(v)
    except Exception: return None


def http(method, path, body=None, prefer=None):
    data = json.dumps(body, ensure_ascii=False).encode('utf-8') if body is not None else None
    h = dict(HDRS)
    if body is not None: h['Content-Type'] = 'application/json'
    if prefer: h['Prefer'] = prefer
    req = urllib.request.Request(URL + path, data=data, method=method, headers=h)
    with urllib.request.urlopen(req, timeout=60) as r:
        t = r.read().decode('utf-8'); return json.loads(t) if t else None


def statut_of(exp, alerte, qte):
    if qte is not None and qte <= 0: return 'archive'
    if exp and exp < TODAY: return 'expire'
    if alerte and alerte <= TODAY: return 'a_renouveler'
    if exp:
        try:
            d = (datetime.date.fromisoformat(exp) - datetime.date.today()).days
            if d <= 30: return 'a_renouveler'
        except Exception: pass
    return 'valide'


def parse_sheet(ws):
    rowsv = []
    for row in ws.iter_rows(min_row=1, max_row=400, max_col=13, values_only=True):
        row = list(row)
        if len(row) < 13: row += [None] * (13 - len(row))
        rowsv.append(row)
    # header = ligne contenant "date de réception"
    hr = None
    for i, row in enumerate(rowsv):
        if any('reception' in norm(c) and 'date' in norm(c) for c in row):
            hr = i; break
    if hr is None:
        return None
    # nom + code dans les lignes au-dessus de l'entête
    nom = code = None
    for row in rowsv[:hr]:
        for j, c in enumerate(row):
            n = norm(c)
            if n.startswith('produit') and not nom:
                for c2 in row[j + 1:]:
                    if s(c2): nom = s(c2); break
            if 'code produit' in n and not code:
                m = re.search(r'code produit\s*:?\s*(.+)', str(c), re.I)
                if m and m.group(1).strip(): code = m.group(1).strip()
                else:
                    for c2 in row[j + 1:]:
                        if s(c2): code = s(c2); break
    lots = []
    cur = None
    for row in rowsv[hr + 1:]:
        drec = d10(row[1]); dsor = d10(row[8])
        inv = numof(row[10])
        if drec:  # nouvelle réception = nouveau lot
            cur = {
                'nom': nom, 'code_produit': code, 'reference': code,
                'n_commande_fournisseur': s(row[2]), 'date_reception': drec, 'date_ouverture': drec,
                'qte_initiale': numof(row[3]), 'lieu_utilisation': s(row[4]), 'emplacement': s(row[4]),
                'date_expiration': d10(row[5]), 'delai_appro': s(row[6]), 'date_alerte': d10(row[7]),
                'quantite': inv if inv is not None else numof(row[3]), 'unite': None,
                'sorties': [],
            }
            lots.append(cur)
        elif dsor and cur is not None:  # sortie du lot courant
            cur['sorties'].append({'date_sortie': dsor, 'qte_sortie': numof(row[9]), 'inventaire_apres': inv, 'n_commande_client': s(row[11])})
            if inv is not None: cur['quantite'] = inv
    return [l for l in lots if l['nom']]


SKIP = re.compile(r'^(trame|sheet2|feuil1|stock|suivi|sortie des produits)', re.I)


def main():
    apply = '--apply' in sys.argv
    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    all_lots = []
    skipped = []
    for sn in wb.sheetnames:
        if SKIP.match(sn.strip()):
            skipped.append(sn); continue
        try:
            lots = parse_sheet(wb[sn])
        except Exception as e:
            skipped.append(sn + ' (err %s)' % e); continue
        if not lots:
            skipped.append(sn); continue
        all_lots.extend(lots)
    nsorties = sum(len(l['sorties']) for l in all_lots)
    print('Feuilles produits :', len(set(l['nom'] for l in all_lots)), '| lots :', len(all_lots), '| sorties :', nsorties)
    print('Feuilles ignorées (', len(skipped), ') :', ', '.join(skipped[:12]), '…' if len(skipped) > 12 else '')
    print('--- échantillon lots ---')
    for l in all_lots[:6]:
        print('  ', str(l['nom'])[:20], '| code', l['code_produit'], '| recep', l['date_reception'], '| pér', l['date_expiration'], '| inv', l['quantite'], '| sorties', len(l['sorties']))
    if not apply:
        print('\nDRY-RUN. --apply pour insérer.'); return
    existing = http('GET', '/rest/v1/produits_perissables?select=code_produit,date_reception')
    seen = set((e.get('code_produit'), e.get('date_reception')) for e in existing)
    ins_prod = ins_mvt = 0
    for l in all_lots:
        key = (l['code_produit'], l['date_reception'])
        if key in seen: continue
        seen.add(key)
        st = statut_of(l['date_expiration'], l['date_alerte'], l['quantite'])
        payload = {k: l[k] for k in ('nom', 'code_produit', 'reference', 'n_commande_fournisseur', 'date_reception', 'date_ouverture', 'qte_initiale', 'lieu_utilisation', 'emplacement', 'date_expiration', 'delai_appro', 'date_alerte', 'quantite', 'unite')}
        payload['statut'] = st
        try:
            res = http('POST', '/rest/v1/produits_perissables', payload, prefer='return=representation')
        except urllib.error.HTTPError as e:
            print('ERR produit', l['nom'], e.code, e.read().decode()[:160]); continue
        pid = res[0]['id'] if res else None
        ins_prod += 1
        if pid and l['sorties']:
            mv = [{'perissable_id': pid, **m} for m in l['sorties']]
            try:
                http('POST', '/rest/v1/mouvements_perissables', mv, prefer='return=minimal')
                ins_mvt += len(mv)
            except urllib.error.HTTPError as e:
                print('ERR mvt', l['nom'], e.code, e.read().decode()[:160])
    print('Inséré : lots', ins_prod, '| mouvements', ins_mvt)


if __name__ == '__main__':
    main()
