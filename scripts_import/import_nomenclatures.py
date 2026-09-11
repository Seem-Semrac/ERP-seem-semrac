# -*- coding: utf-8 -*-
# Import nomenclatures SEEM/SEMRAC depuis les TCD VPU.
# DRY-RUN par defaut. Ajouter --apply pour ecrire reellement.
# Regles : temps en milliemes d'heures stockes TELS QUELS ; reglage -> temps fixe ;
#          doublons (ref,phase) -> MAX ; homme/machine via process.requiert_machine (sinon homme) ;
#          statut nomenclature = en_cours ; libelles non reconnus conserves sans process_id ;
#          ignorer lignes Total/Sous-total + pointeurs 'VOIR GAMME SUIVANTE'.
import openpyxl, re, sys, json, unicodedata, urllib.request, urllib.error

SUPABASE_URL = 'https://vyqgrasezpyqjwvijwvv.supabase.co'
ANON = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InZ5cWdyYXNlenB5cWp3dmlqd3Z2Iiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzgwNTE4MTcsImV4cCI6MjA5MzYyNzgxN30.Optt1QD-Ai3Xm8mk4f8fsvKry9-L2SlTBNB9KgVG3ls'
DRY_RUN = '--apply' not in sys.argv

FILES = [
    (r'C:\Users\eddys\Desktop\vpux_20260622_105849.xlsx', 'Seem'),
    (r'C:\Users\eddys\Desktop\vpux_20260622_132039.xlsx', 'Semrac'),
]

def norm(s):
    s = unicodedata.normalize('NFD', str(s or '')).encode('ascii', 'ignore').decode().upper()
    return re.sub(r'[^A-Z0-9 ]', ' ', s).strip()

def rest(method, path, body=None, params=''):
    url = SUPABASE_URL + '/rest/v1/' + path + params
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, method=method)
    req.add_header('apikey', ANON); req.add_header('Authorization', 'Bearer ' + ANON)
    req.add_header('Content-Type', 'application/json')
    if method in ('POST', 'PATCH'): req.add_header('Prefer', 'return=minimal')
    with urllib.request.urlopen(req, timeout=30) as r:
        txt = r.read().decode()
        return json.loads(txt) if txt else []

# --- referentiel process pour matching flou ---
procs = rest('GET', 'process_atelier', None, '?select=id,nom,activite,requiert_machine,machine_id,poste_id')
STOP = set(['SEEM', 'SEMRAC', 'POSTE', 'ZONE', 'MACHINE', 'MANUEL', 'DE', 'LA', 'LE', 'MC', 'CN', 'OAS'])
def toks(s):
    return set(t for t in norm(s).split() if len(t) >= 3 and t not in STOP)
proc_by_site = {}
for p in procs:
    proc_by_site.setdefault(p.get('activite'), []).append((p['id'], p['nom'], p.get('requiert_machine'), p.get('machine_id'), toks(p['nom'])))

# Synonymes : verbe d'operation (libelle TCD) -> fragment du nom de process. Ordre = priorite (specifique avant generique).
SYN = [
    ('KIA 1', 'KIA 1'), ('KIA 2', 'KIA 2'), ('KIA 3', 'KIA 3'), ('KIA 4', 'KIA 4'), ('KIA', 'KIA'),
    ('TRONCON', 'TRONCONNEUSE'), ('CISAIL', 'CISAILLAGE'), ('POINCON', 'POINCON'),
    ('PLIAGE', 'PLIEUSE'), ('PLI HFB', 'PLIEUSE'), ('PLIEUSE', 'PLIEUSE'),
    ('EBAVUR', 'EBAVURAGE'), ('BROSS', 'BROSSAGE'), ('MEUL', 'MEULAGE'), ('LAMA', 'LAMAGE'),
    ('FRAISUR', 'FRAISURAGE'), ('SOUD', 'SOUDURE'), ('SERIGRAPH', 'SERIGRAPHIE'),
    ('THERMOCOL', 'THERMOCOLLAGE'), ('ASSEMBL', 'ASSEMBLAGE'), ('MONTAGE', 'ASSEMBLAGE'),
    ('OXYD', 'OXYDATION'), ('OAS', 'OXYDATION'), ('DAHLIH', 'DAHLIH'), ('STAMA', 'STAMA'),
    ('DMC', 'DMC'), ('SERIG', 'SERIGRAPHIE'), ('EMBALL', 'EMBALLAGE'), ('DEBALL', 'EMBALLAGE'),
    ('METROLOG', 'CONTROLE'), ('CONTROLE', 'CONTROLE'), ('VERIF', 'CONTROLE'),
]
def match_process(libelle, site):
    L = norm(libelle)
    cands = proc_by_site.get(site, []) + proc_by_site.get('both', [])
    for kw, frag in SYN:
        if kw in L:
            for (pid, pnom, rm, pmid, ptok) in cands:
                if frag in norm(pnom):
                    return (pid, pnom, rm, pmid)
    Ltok = set(t for t in L.split() if len(t) >= 3)
    best = None; bestscore = 0
    for (pid, pnom, rm, pmid, ptok) in cands:
        score = len(Ltok & ptok) + sum(1 for t in ptok if t in L)
        if score > bestscore: bestscore = score; best = (pid, pnom, rm, pmid)
    return best if bestscore >= 1 else None

# --- nomenclatures existantes (anti-doublon) ---
def tout(path, select):
    # Lecture PAGINEE : PostgREST plafonne chaque reponse (1000 lignes par defaut).
    out, off = [], 0
    while True:
        page = rest('GET', path, None, '?select=' + select + '&order=id&limit=1000&offset=' + str(off))
        out += page
        if len(page) < 1000: return out
        off += 1000
existing = tout('nomenclatures', 'id,entite,code_ref_produit,statut')
exmap = {(n.get('entite'), str(n.get('code_ref_produit') or '').strip()): n['id'] for n in existing}
# Une nomenclature SUIVIE au journal EN 9100 (validee, ou validee puis devalidee, ou revision d'une
# validee) n'est JAMAIS reecrite par l'import : il la devaliderait et remplacerait sa gamme en
# passant par la base directement, donc hors du journal. Nouvel indice via l'ERP.
try:
    suivies_ids = {j.get('nomenclature_id') for j in tout('nomenclature_journal', 'nomenclature_id')}
except urllib.error.HTTPError as e:
    if e.code != 404: raise           # table absente (cloud avant cloud-5) : seules les validees comptent
    suivies_ids = set()
valides = {(n.get('entite'), str(n.get('code_ref_produit') or '').strip()) for n in existing if n.get('statut') == 'valide' or n.get('id') in suivies_ids}

rep = {'create': 0, 'update': 0, 'ignored': 0, 'errors': [], 'unmatched': {}, 'samples': [], 'by_site': {}}

for path, site in FILES:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Feuil1']
    hi = None; hdr = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        vals = [str(c).strip() if c else '' for c in row]
        if 'Code Art.' in vals: hdr = vals; hi = i; break
    col = {}
    for j, h in enumerate(hdr):
        hl = h.lower()
        if h == 'Code Art.': col['ref'] = j
        elif h == 'Designation': col['des'] = j
        elif h == 'Libelle Phase': col['lib'] = j
        elif h == 'Phase': col['ph'] = j
        elif 'millieme' in hl: col['t'] = j
        elif 'fixe/variable' in hl: col['fv'] = j
    if 'fv' not in col:
        for j, h in enumerate(hdr):
            if h == 'T': col['fv'] = j; break
    noms = {}
    for row in list(ws.iter_rows(values_only=True))[hi + 1:]:
        ref = row[col['ref']] if col.get('ref') is not None else None
        if ref is None or str(ref).strip() == '': continue
        if isinstance(ref, float) and ref.is_integer(): ref = int(ref)
        refs_ = str(ref).strip()
        lib = str(row[col['lib']]).strip() if col.get('lib') is not None and row[col['lib']] else ''
        des = str(row[col['des']]).strip() if col.get('des') is not None and row[col['des']] else ''
        nl = norm(lib + ' ' + des)
        if re.search(r'\b(TOTAL|SOUS TOTAL|TOTAL GENERAL)\b', nl): rep['ignored'] += 1; continue
        if 'VOIR GAMME SUIVANTE' in norm(lib): rep['ignored'] += 1; continue
        fvraw = str(row[col['fv']]).strip().upper() if col.get('fv') is not None and row[col['fv']] else ''
        is_regl = bool(re.search(r'R[E]GLAGES?', norm(lib)))
        est_fixe = (fvraw == 'F') or is_regl
        t = row[col['t']] if col.get('t') is not None else None
        try: tval = float(t) if t is not None else None
        except Exception: tval = None
        if tval is None: rep['ignored'] += 1; continue
        phk = (str(row[col['ph']]).strip().lower() if col.get('ph') is not None and row[col['ph']] else lib.lower())
        d = noms.setdefault(refs_, {'desig': des, 'phases': {}})
        if des and not d['desig']: d['desig'] = des
        cur = d['phases'].get(phk)
        if cur is None or tval > cur['t']:
            d['phases'][phk] = {'lib': lib, 'ph': str(row[col['ph']] or '').strip(), 'fixe': est_fixe, 't': tval}
    rep['by_site'][site] = len(noms)
    for refs_, d in noms.items():
        etapes = []; ordre = 0; pending_regl = 0.0
        for ph in d['phases'].values():            # ordre d'apparition = ordre de gamme
            if re.search(r'REGLAGES?', norm(ph['lib'])):   # REGLAGE -> rattache au process AVAL (pas d'etape distincte)
                pending_regl += ph['t']; rep['reglage_merged'] = rep.get('reglage_merged', 0) + 1; continue
            ordre += 1
            m = match_process(ph['lib'], site)
            ressource = 'machine' if (m and m[2]) else 'homme'
            if not m: rep['unmatched'][ph['lib']] = rep['unmatched'].get(ph['lib'], 0) + 1
            etapes.append({
                'ordre': ordre, 'phase': ph['ph'], 'nom': ph['lib'],
                'process_id': m[0] if m else None, 'process_nom': m[1] if m else None,
                'machine_id': (m[3] if (m and m[2]) else None),   # taux machine reel (cout_h + OPEX poste) resolu via machine_id
                'est_fixe': ph['fixe'], 'ressource': ressource,
                'temps_variable_mille': ph['t'],          # millieme BRUT (op) ; heures = /1000
                'temps_reglage_mille': pending_regl,      # reglage(s) en amont rattache(s) ici (setup, par lot)
                'site': site, 'source': 'import_tcd',
            })
            pending_regl = 0.0
        if pending_regl > 0 and etapes:                # reglage en fin de gamme -> dernier process
            etapes[-1]['temps_reglage_mille'] = (etapes[-1]['temps_reglage_mille'] or 0) + pending_regl
        elif pending_regl > 0:                          # gamme uniquement reglage -> 1 etape porteuse
            etapes.append({'ordre': 1, 'phase': '', 'nom': 'Reglage', 'process_id': None, 'process_nom': None,
                           'est_fixe': True, 'ressource': 'homme', 'temps_variable_mille': 0,
                           'temps_reglage_mille': pending_regl, 'site': site, 'source': 'import_tcd'})
        payload = {'entite': site, 'num_nom': refs_, 'code_ref_produit': refs_,
                   'description': d['desig'] or None, 'statut': 'en_cours', 'indice': 'A',
                   'type_nom': 'standard', 'etapes_production': etapes}
        key = (site, refs_)
        if key in valides:
            rep['valide_ignoree'] = rep.get('valide_ignoree', 0) + 1
            continue
        if key in exmap:
            rep['update'] += 1
            if not DRY_RUN:
                try: rest('PATCH', 'nomenclatures', payload, '?id=eq.' + exmap[key])
                except Exception as e: rep['errors'].append(refs_ + ': ' + str(e))
        else:
            rep['create'] += 1
            if not DRY_RUN:
                try: rest('POST', 'nomenclatures', payload)
                except Exception as e: rep['errors'].append(refs_ + ': ' + str(e))
        if len(rep['samples']) < 3:
            rep['samples'].append((site, refs_, d['desig'], etapes[:4]))

print('================ RAPPORT IMPORT NOMENCLATURES ================')
print('MODE :', 'DRY-RUN (aucune ecriture)' if DRY_RUN else '*** APPLY (ecriture reelle) ***')
print('Nomenclatures par site :', rep['by_site'])
print('Suivies EN 9100 ignorees (validees ou deja au journal, jamais reecrites) :', rep.get('valide_ignoree', 0))
print('Creations :', rep['create'], '| Mises a jour :', rep['update'], '| Lignes ignorees :', rep['ignored'], '| Reglages fusionnes au process aval :', rep.get('reglage_merged', 0))
print('Erreurs :', len(rep['errors']))
for e in rep['errors'][:10]: print('   ERR', e)
um = sorted(rep['unmatched'].items(), key=lambda x: -x[1])
print('Libelles NON rattaches a un process (%d distincts) — top 25 :' % len(um))
for l, n in um[:25]: print('   - %-30s x%d' % (l[:30], n))
print('--- Exemples de gammes construites ---')
for site, ref, des, et in rep['samples']:
    print('  [%s] %s — %s (%d etapes montrees)' % (site, ref, des, len(et)))
    for e in et:
        print('      #%d %-22s var_mille=%s regl_mille=%s fixe=%s ress=%s proc=%s' % (
            e['ordre'], e['nom'][:22], e['temps_variable_mille'], e['temps_reglage_mille'], e['est_fixe'], e['ressource'], e['process_nom']))
print('=============================================================')
if DRY_RUN: print('>> Relancer avec --apply pour ecrire en base.')
